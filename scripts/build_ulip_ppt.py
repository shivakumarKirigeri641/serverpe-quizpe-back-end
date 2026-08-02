"""Build the ULIP use-case submission deck for ServerPe App Solutions.

Reads the official ULIP template (src/temp/Updated Use Cases Format 1.pptx),
fills the intro slides and clones the sample use-case slide once per use case.
"""

import copy
from pptx import Presentation
from pptx.util import Pt, Emu

TEMP = "src/temp"
SRC = f"{TEMP}/Updated Use Cases Format 1.pptx"
OUT = f"{TEMP}/ServerPe_ULIP_UseCases.pptx"
LOGO = f"{TEMP}/ServerPe_Logo.png"

SUBMIT_DATE = "20 July 2026"
CONTACT = "Shivakumar Kirigeri"
EMAIL = "admin@serverpe.in"
MOBILE = "+91 98861 22415"


# ---------------------------------------------------------------- helpers
def set_text(shape, lines, size=None, bullet=False):
    """Replace a shape's text, keeping the formatting of its first run."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    if not p0.runs:
        p0.add_run()
    proto_r = copy.deepcopy(p0.runs[0]._r)
    proto_p = copy.deepcopy(p0._p)

    # wipe everything after the first paragraph
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)

    def fill(par, text):
        while len(par.runs) > 1:
            par.runs[-1]._r.getparent().remove(par.runs[-1]._r)
        if not par.runs:
            par._p.append(copy.deepcopy(proto_r))
        run = par.runs[0]
        run.text = ("• " + text) if bullet else text
        if size:
            run.font.size = Pt(size)

    fill(p0, lines[0])
    for line in lines[1:]:
        newp = copy.deepcopy(proto_p)
        p0._p.getparent().append(newp)
        par = tf.paragraphs[-1]
        fill(par, line)
    tf.word_wrap = True


def by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(name)


def clone_slide(prs, src_slide):
    """Append a copy of src_slide (same layout, all shapes) to the deck."""
    dest = prs.slides.add_slide(src_slide.slide_layout)
    for sh in list(dest.shapes):
        sh._element.getparent().remove(sh._element)
    for sh in src_slide.shapes:
        dest.shapes._spTree.append(copy.deepcopy(sh._element))
    return dest


# ---------------------------------------------------------------- content
USE_CASES = [
    {
        "title": "Use Case 1 : ChallanAlerts - Proactive Vehicle Compliance "
                 "Monitoring for Fleets & Logistics Operators",
        "problem": [
            "Fleet owners, transporters and 3PL operators track compliance for "
            "hundreds of vehicles across separate portals - RC/fitness, "
            "insurance, PUC, permit, tax, FASTag and eChallan.",
            "Checks are manual, periodic and error-prone. A single missed "
            "renewal causes check-post detention, penalties, off-road days, "
            "delayed consignments and uninsured-liability exposure.",
            "No consolidated, real-time view of fleet-level compliance health "
            "exists for the operator or for the consignor auditing the fleet.",
        ],
        "solution": [
            "ChallanAlerts - a consent-based SaaS + mobile platform giving a "
            "single compliance control tower for the entire fleet.",
            "Modules using ULIP data: (1) Vehicle onboarding & consent manager, "
            "(2) Compliance Sync - pulls RC, fitness, tax, permit, insurance, "
            "PUC, FASTag and eChallan records via ULIP, (3) Rule & Alert Engine "
            "- T-30/15/7/1 day reminders over WhatsApp, SMS, email and push, "
            "(4) Fleet dashboard with vehicle-wise compliance score and "
            "exception list, (5) Audit-ready reports for shippers and auditors.",
            "Benefit: near-zero missed renewals, fewer detentions and fines, "
            "higher vehicle uptime and a verifiable compliance trail.",
        ],
        "additional": [
            "Predictive compliance-risk scoring per vehicle and per route.",
            "Driver licence validity tracking via Sarathi.",
            "Trip-level compliance using FASTag transaction trails and e-Way "
            "Bill linkage.",
            "Plug-in APIs so TMS/ERP and insurers can consume the compliance "
            "status directly.",
        ],
        "targets": [
            "Logistics & transport fleet operators, 3PL and freight aggregators",
            "Small truck owners and transport associations",
            "Corporate, staff-transport and cab-aggregator fleets",
            "Leasing, rental and vehicle-finance companies",
            "Not captive - offered commercially as a B2B / B2C SaaS product",
        ],
        "integrations": "Integrations: ULIP VAHAN (RC / fitness / tax), "
                        "eChallan, Insurance (IIB), PUC, National Permit, "
                        "FASTag-NETC, Sarathi (DL)",
    },
    {
        "title": "Use Case 2 : VerifyVahan - Pre-Transaction Verification of "
                 "Pre-Owned Commercial & Passenger Vehicles",
        "problem": [
            "The pre-owned vehicle market - especially commercial vehicles used "
            "in logistics - is opaque. A buyer cannot independently confirm "
            "registration validity, hypothecation, permit, fitness, insurance, "
            "PUC, blacklist status or pending challan dues.",
            "This leads to fraud, disputed ownership transfers, vehicles "
            "purchased with unpaid liabilities, and rejected finance cases.",
            "Dealers, financiers and fleet buyers have no fast, standardised "
            "due-diligence check before money changes hands.",
        ],
        "solution": [
            "VerifyVahan - an instant, consent-based vehicle verification "
            "report generated from authorised ULIP data sources.",
            "Modules using ULIP data: (1) Single-search verification - RC, "
            "owner-serial, fitness, tax and permit status, (2) Encumbrance & "
            "risk checks - hypothecation/financier, blacklist, pending "
            "eChallan, (3) Insurance and PUC validity, (4) Risk score with "
            "red-flag summary and shareable PDF report, (5) Bulk verification "
            "API for dealers, auction platforms and lenders.",
            "Benefit: transparent transactions, informed pricing, lower fraud "
            "and dispute rates, faster ownership transfer and faster loan "
            "decisions.",
        ],
        "additional": [
            "Valuation assist using vehicle age, permit class and compliance "
            "history.",
            "Financier NOC and RTO transfer checklist tracking.",
            "Fraud-pattern analytics across repeat sellers and vehicles.",
            "Direct integration into lender loan-origination systems and "
            "fleet-acquisition due diligence workflows.",
        ],
        "targets": [
            "Used-vehicle buyers and sellers (commercial & passenger)",
            "Vehicle dealers, brokers and online auction platforms",
            "NBFCs and banks in vehicle finance; insurers",
            "Logistics companies acquiring or leasing second-hand fleet",
            "Not captive - offered commercially to businesses and consumers",
        ],
        "integrations": "Integrations: ULIP VAHAN (RC / owner / hypothecation / "
                        "fitness / blacklist), eChallan, Insurance (IIB), PUC, "
                        "National Permit",
    },
    {
        "title": "Use Case 3 : FASTag Trip & Toll Intelligence - End-of-Day "
                 "Toll Movement Digest for Vehicle Owners and Fleets",
        "problem": [
            "A vehicle owner or fleet operator has no consolidated daily view "
            "of where their vehicles actually travelled. FASTag debits are "
            "visible only as bank SMS lines with no plaza, route or trip "
            "context.",
            "Toll spend cannot be reconciled against trips, so leakage - "
            "unauthorised vehicle use, off-route running, duplicate or "
            "erroneous debits - goes undetected.",
            "Unexpected low-balance blocks at a plaza stop the vehicle and "
            "delay the consignment; owners learn of it only after the halt.",
        ],
        "solution": [
            "A FASTag Trip & Toll Intelligence module inside ChallanAlerts, "
            "delivering an automated end-of-day digest per owner and per fleet.",
            "Modules using ULIP data: (1) Recent transaction feed - plaza-wise "
            "FASTag debits with timestamp, amount and lane/plaza name, "
            "(2) Toll-crossed trail - sequence of plazas passed during the day, "
            "reconstructed into a route view, (3) End-of-day digest pushed at a "
            "fixed cut-off over WhatsApp, email and app - vehicles moved, "
            "plazas crossed, total toll spend, exceptions, (4) Tag health - "
            "low balance, blacklisted/exception tag alerts, (5) Owner-wise and "
            "vehicle-wise toll expense reports for accounting and GST records.",
            "Benefit: daily visibility of vehicle movement, verifiable toll "
            "spend, early detection of misuse and no surprise plaza halts.",
        ],
        "additional": [
            "Route inference and ETA estimation from consecutive plaza "
            "crossings.",
            "Trip reconciliation against e-Way Bill and consignment records.",
            "Toll-cost benchmarking per route and per vehicle class for freight "
            "pricing.",
            "Geo-fence and unusual-movement alerts; auto top-up triggers on low "
            "tag balance.",
        ],
        "targets": [
            "Fleet and transport operators reconciling daily toll spend",
            "Individual commercial-vehicle owners with drivers on the road",
            "3PL, freight forwarders and consignors tracking shipment movement",
            "Corporate fleets, rental and leasing companies",
            "Not captive - offered commercially as a B2B / B2C SaaS module",
        ],
        "integrations": "Integrations: ULIP FASTag-NETC (transaction history / "
                        "toll plaza crossings / tag status & balance), VAHAN "
                        "(vehicle-owner mapping), e-Way Bill (trip "
                        "reconciliation)",
    },
]


# ---------------------------------------------------------------- build
prs = Presentation(SRC)
s1, s2, sample = prs.slides[0], prs.slides[1], prs.slides[2]

# --- Slide 1: title
set_text(by_name(s1, "TextBox 2"),
         ["Unified Logistics Interface Platform (ULIP)",
          "Submission of Use-Cases",
          "ServerPe App Solutions"])
set_text(by_name(s1, "TextBox 7"),
         [f"Date of Submission: {SUBMIT_DATE}",
          f"Person of Contact: {CONTACT}",
          f"Email ID: {EMAIL}",
          f"Mobile Number: {MOBILE}"])
s1.shapes.add_picture(LOGO, Emu(9700000), Emu(600000), height=Emu(1000000))

# --- Slide 2: introduction
set_text(by_name(s2, "TextBox 4"), [
    "A. Name & corporate address of the Company:",
    "ServerPe App Solutions, The Orchard, HMT Watch Factory Main Road, "
    "Jalahalli, Bangalore - 560013, Karnataka, India.  Website: "
    "https://serverpe.in",
    "",
    "B. Company Incorporated in (Month/Year):  <MONTH / YEAR>",
    "",
    "C. Brief about the Company (150 words max):",
    "ServerPe App Solutions is an Indian technology company building secure, "
    "scalable digital platforms for the mobility and logistics sector. We "
    "develop web and mobile applications, SaaS platforms, API integrations and "
    "compliance-focused products that simplify operations for individuals and "
    "enterprises. Our vision is to be a trusted technology partner delivering "
    "platforms that improve accessibility, compliance and operational "
    "efficiency across mobility and related sectors. Our current products - "
    "ChallanAlerts and VerifyVahan - address vehicle compliance monitoring and "
    "pre-transaction vehicle verification. We follow a privacy-first, "
    "consent-based approach: data obtained from authorised APIs is processed "
    "securely, used only for approved purposes, protected by access controls "
    "and never shared with unauthorised third parties.",
    "",
    "D. Services offered by the Company (150 words max):",
    "Vehicle compliance solutions; mobility technology platforms; API "
    "integration services; SaaS platform development; web and mobile "
    "application development; cloud solutions; and business process "
    "automation. Our technology stack comprises REST APIs, Node.js, React, "
    "PostgreSQL, secure authentication, cloud infrastructure, multi-channel "
    "notification systems and consent-based data processing.",
    "",
    "E. No. of Clients currently using Company's services:  <NUMBER>",
    "",
    "F. Type of users currently using Company's services:",
    "Individual vehicle owners, fleet and transport operators, vehicle dealers "
    "and buyers of pre-owned vehicles, and enterprises managing corporate "
    "fleets.",
    "",
    "G. Net Worth of the Company:  Rs. <AMOUNT>",
], size=11)

# --- Use-case slides (clone the sample, drop the (SAMPLE) tag)
for uc in USE_CASES:
    sl = clone_slide(prs, sample)
    by_name(sl, "TextBox 29")._element.getparent().remove(
        by_name(sl, "TextBox 29")._element)
    set_text(by_name(sl, "TextBox 24"), [uc["title"]], size=18)
    set_text(by_name(sl, "TextBox 30"), uc["problem"], size=10, bullet=True)
    set_text(by_name(sl, "TextBox 3"), uc["solution"], size=10, bullet=True)
    set_text(by_name(sl, "TextBox 33"), uc["additional"], size=10, bullet=True)
    set_text(by_name(sl, "TextBox 34"), uc["targets"], size=10, bullet=True)
    set_text(by_name(sl, "Rectangle 31"), [uc["integrations"]], size=11)
    set_text(by_name(sl, "TextBox 35"), [""])

# drop the original sample slide from the deck
xml_slides = prs.slides._sldIdLst
xml_slides.remove(list(xml_slides)[2])

prs.save(OUT)
print("written:", OUT, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
